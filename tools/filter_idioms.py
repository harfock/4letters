#!/usr/bin/env python3
"""Build a clean four-character idiom candidate list from the MOE workbook.

Usage:
  python tools/filter_idioms.py dict_idioms_2020_20260929.xlsx data/idioms.json

The output intentionally contains headwords and pronunciation only; it does not
redistribute the dictionary's definitions, examples, or literary citations.
"""
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

REDIRECT = re.compile(r'^\s*即[「『《](.+?)[」』》][。；;，,]?\s*見[「『《](.+?)[」』》]條')


def main():
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    source, destination = map(Path, sys.argv[1:])
    wb = load_workbook(source, read_only=True, data_only=True)
    ws = wb['成語資料']
    records = []
    skipped_redirects = []
    seen = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        word = row[1]
        pinyin = row[3]
        definition = row[4] or ''
        marker = str(row[21] or '').strip()
        if marker != '主條成語' or not isinstance(word, str):
            continue
        word = word.strip()
        if len(word) != 4 or word in seen:
            continue
        match = REDIRECT.search(str(definition))
        if match and match.group(1) == match.group(2) and match.group(1) != word:
            skipped_redirects.append({'word': word, 'canonical': match.group(1)})
            continue
        seen.add(word)
        records.append({'word': word, 'pinyin': str(pinyin or '').strip()})
    records.sort(key=lambda item: item['word'])
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(records, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f'Candidates: {len(records)}')
    print(f'Explicit redirects skipped: {len(skipped_redirects)}')
    print(f'Output: {destination}')
    if '紙醉金迷' not in seen or '金迷紙醉' in seen:
        raise SystemExit('Canonical-entry validation failed for 紙醉金迷 / 金迷紙醉')


if __name__ == '__main__':
    main()
