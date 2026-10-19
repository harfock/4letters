#!/usr/bin/env python3
"""Generate 21 static connected crossword levels from an MOE workbook.
Usage: python tools/build_levels.py source.xlsx public/data/levels.json
Only headwords and pinyin are exported, never official definitions or examples.
"""
import collections, json, random, re, sys
from pathlib import Path
from openpyxl import load_workbook
REDIRECT=re.compile(r'^\s*即[「『《](.+?)[」』》][。；;，,]?\s*見[「『《](.+?)[」』》]條')

def read_words(path):
    wb=load_workbook(path,read_only=True,data_only=True); ws=wb['成語資料']; result=[]; seen=set()
    for row in ws.iter_rows(min_row=2,values_only=True):
        word,py=row[1],row[3]; definition=str(row[4] or ''); marker=str(row[21] or '').strip()
        if not isinstance(word,str): continue
        word=word.strip()
        if len(word)!=4: continue
        m=REDIRECT.search(definition)
        if m and m.group(1)==m.group(2) and m.group(1)!=word: continue
        if marker!='主條成語' or word in seen: continue
        seen.add(word); result.append({'word':word,'pinyin':str(py or '').strip()})
    if '紙醉金迷' not in seen or '金迷紙醉' in seen: raise ValueError('Canonical-entry validation failed')
    return sorted(result,key=lambda x:x['word'])

def cells(p):
    return {(p['row']+(i if p['direction']=='down' else 0),p['col']+(i if p['direction']=='across' else 0)):ch for i,ch in enumerate(p['word'])}

def make_level(pool,n,level):
    for attempt in range(70):
        candidates=list(pool); rng=random.Random(level*1009+attempt); rng.shuffle(candidates)
        if not candidates: return None
        ps=[{'word':candidates[0]['word'],'pinyin':candidates[0]['pinyin'],'row':0,'col':0,'direction':'across','initialReveal':True,'clue':''}]
        used={candidates[0]['word']}; remaining=[x for x in candidates[1:] if x['word'] not in used]
        while len(ps)<n:
            occ=collections.defaultdict(list)
            cross=collections.defaultdict(list)
            for pi,p in enumerate(ps):
                for pos,ch in cells(p).items():
                    occ[pos].append((pi,ch,p['direction']))
                    cross[(ch,p['direction'])].append(pos)
            choices=[]
            for item in remaining:
                for d in ('across','down'):
                    other='down' if d=='across' else 'across'
                    for off,ch in enumerate(item['word']):
                        for r,c in cross.get((ch,other),[]):
                            p={'word':item['word'],'pinyin':item['pinyin'],'row':r-(off if d=='down' else 0),'col':c-(off if d=='across' else 0),'direction':d,'initialReveal':False,'clue':''}
                            nc=cells(p); valid=True; crossings=0
                            for pos,cc in nc.items():
                                old=occ.get(pos,[])
                                if old:
                                    if any(ec!=cc or ed==d for _,ec,ed in old): valid=False; break
                                    crossings+=1
                                else:
                                    rr,cl=pos
                                    if any(a in occ and a not in nc for a in ((rr-1,cl),(rr+1,cl),(rr,cl-1),(rr,cl+1))): valid=False; break
                            if valid and crossings:
                                allpos=list(occ)+list(nc); h=max(x[0] for x in allpos)-min(x[0] for x in allpos)+1; w=max(x[1] for x in allpos)-min(x[1] for x in allpos)+1
                                choices.append((crossings*10-abs(w-h)*.02-w*h*.0005+rng.random(),p,item))
            if not choices: break
            _,p,item=max(choices,key=lambda x:x[0])
            ps.append(p); used.add(item['word']); remaining=[x for x in remaining if x['word']!=item['word']]
        if len(ps)==n:
            mr=min(p['row'] for p in ps); mc=min(p['col'] for p in ps)
            for p in ps:p['row']-=mr;p['col']-=mc
            for i,p in enumerate(ps):p['initialReveal']=(1+((i+level)%2)) if i<n//2 else 0
            return ps
    return None

def validate(levels):
    allwords=set()
    for level in levels:
        ps=level['placements']; occ=collections.defaultdict(list); adj=[set() for _ in ps]
        for i,p in enumerate(ps):
            if p['word'] in allwords: raise ValueError('Duplicate word: '+p['word'])
            allwords.add(p['word'])
            for pos,ch in cells(p).items():occ[pos].append((i,ch,p['direction']))
        for pos,es in occ.items():
            if len({e[1] for e in es})>1:raise ValueError(f'Character conflict at {pos}')
            if len(es)>1:
                if len({e[2] for e in es})<2:raise ValueError(f'Same-direction overlap at {pos}')
                for i,_,_ in es:
                    for j,_,_ in es:
                        if i!=j:adj[i].add(j)
        seen={0}; q=collections.deque([0])
        while q:
            i=q.popleft()
            for j in adj[i]-seen:seen.add(j);q.append(j)
        if len(seen)!=len(ps):raise ValueError(f'Level {level["level"]} is disconnected')
    if len(allwords)!=420:raise ValueError(f'Expected 420 unique words, got {len(allwords)}')

def main():
    if len(sys.argv)!=3:raise SystemExit('Usage: python tools/build_levels.py source.xlsx public/data/levels.json')
    unused=read_words(sys.argv[1]); levels=[]
    for lev,n in enumerate(range(10,31),1):
        ps=None
        for size in (1000,1300,1600):
            ps=make_level(unused[:size],n,lev)
            if ps:break
        if not ps:raise SystemExit(f'Could not build connected level {lev}; no data emitted')
        selected={p['word'] for p in ps}; unused=[w for w in unused if w['word'] not in selected]
        levels.append({'level':lev,'placements':ps})
    validate(levels); out={'schemaVersion':1,'sourceVersion':'dict_idioms_2020_20260929','levels':levels}
    dest=Path(sys.argv[2]);dest.parent.mkdir(parents=True,exist_ok=True);dest.write_text(json.dumps(out,ensure_ascii=False,separators=(',',':'))+'\n',encoding='utf-8')
    print(f'PASS: 21 levels, 420 unique connected placements -> {dest}')
if __name__=='__main__':main()
